import io, json, os, time
from pathlib import Path

import streamlit as st
from pypdf import PdfReader
from pptx import Presentation
from google import genai
from google.genai import types

st.set_page_config(page_title="AI 학습 도우미 v2", page_icon="📚", layout="wide")
st.title("📚 AI 학습 도우미 v2")
st.caption("텍스트 + 그림/도표를 함께 분석하여 시험 대비 자료와 문제를 만드는 서비스")

# ---------- extraction ----------
def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append({"page": i, "text": page.extract_text() or ""})
    return pages

def pptx_to_images(data):
    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        images = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.shape_type == 13:  # picture
                try:
                    blob = shape.image.blob
                    mime = "image/png" if shape.image.ext in ("png",) else "image/jpeg"
                    images.append({
                        "mime": mime,
                        "raw_bytes": blob
                    })
                except Exception:
                    pass
        slides.append({"page": i, "text": "\n".join(texts), "images": images})
    return slides

def extract(uploaded):
    data = uploaded.getvalue()
    ext = Path(uploaded.name).suffix.lower()
    if ext == ".pdf":
        pages = extract_pdf(data)
        for p in pages:
            p["images"] = []
        return pages
    if ext == ".pptx":
        return pptx_to_images(data)
    raise ValueError("PDF 또는 PPTX만 지원합니다.")

# ---------- Gemini AI 호출 및 자동 복구 로직 ----------
def ai_analyze(slides, settings):
    # 1. API 키 불러오기 (Streamlit Secrets -> os 환경변수 -> 사이드바 입력 순서)
    api_key = None
    if "GEMINI_API_KEY" in st.secrets:
        api_key = st.secrets["GEMINI_API_KEY"]
    elif os.getenv("GEMINI_API_KEY"):
        api_key = os.getenv("GEMINI_API_KEY")
    elif settings.get("api_key_input"):
        api_key = settings["api_key_input"]

    if not api_key:
        return None, "Gemini API 키가 없습니다. Streamlit Cloud Secrets에 'GEMINI_API_KEY'를 추가해주세요."

    client = genai.Client(api_key=api_key)

    system_instruction = f"""
너는 대학 강의자료 기반 시험 대비 학습 도우미다.

반드시 제공된 강의자료를 근거로 분석한다.
자료에 없는 사실을 임의로 추가하지 않는다.
슬라이드의 그림, 그래프, 표, 도식 등도 적극적으로 읽고 시험 포인트로 추출한다.

[사용자 설정]
문제 수: {settings['n']}
문제 유형: {settings['types']}
난이도: {settings['difficulty']}
출제 스타일: {settings['styles']}

출력은 반드시 JSON 스키마에 맞춰서만 응답해야 한다.
{{
  "study_summary": "전체 요약",
  "key_concepts": [
    {{"concept":"개념명", "explanation":"설명", "importance":"high/medium/low"}}
  ],
  "visual_points": [
    {{"slide":1, "type":"diagram/image/graph/table", "what_it_shows":"내용", "exam_point":"시험 포인트"}}
  ],
  "questions": [
    {{
      "type":"객관식/OX/단답형",
      "source_slide": 1,
      "uses_visual": true,
      "question":"문제 내용",
      "choices":["A. ...","B. ...","C. ...","D. ..."],
      "answer":"정답",
      "explanation":"해설"
    }}
  ]
}}
"""

    contents = []
    for s in slides:
        txt = s.get("text", "")
        contents.append(f"\n\n--- SLIDE/PAGE {s['page']} ---\n{txt[:10000]}")
        for img in s.get("images", []):
            contents.append(
                types.Part.from_bytes(
                    data=img["raw_bytes"],
                    mime_type=img["mime"]
                )
            )

    # 503 과부하 대비: 기본 모델 실패 시 예비 모델 목록으로 자동 순회
    models_to_try = [settings["model"], "gemini-3.6-flash", "gemini-3.8-flash"]
    # 중복 제거
    models_to_try = list(dict.fromkeys(models_to_try))

    last_error = ""

    for target_model in models_to_try:
        # 각 모델당 최대 2회 재시도
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=target_model,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        system_instruction=system_instruction,
                        response_mime_type="application/json",
                        temperature=0.2
                    )
                )
                return json.loads(response.text), None
            except Exception as e:
                last_error = str(e)
                # 503 과부하인 경우 3초 대기 후 재시도
                if "503" in last_error:
                    time.sleep(3)
                    continue
                # 모델 404인 경우 즉시 다음 후보 모델로 변경
                elif "404" in last_error:
                    break
                else:
                    return None, f"오류 발생: {last_error}"

    return None, f"구글 API 서버 과부하로 처리에 실패했습니다. 1~2분 후 다시 눌러주세요. (세부 내용: {last_error})"

# ---------- UI ----------
with st.sidebar:
    st.header("⚙️ 설정")
    
    # Secrets가 없을 경우를 대비한 수동 입력창
    has_secret_key = "GEMINI_API_KEY" in st.secrets or os.getenv("GEMINI_API_KEY")
    if not has_secret_key:
        api_key_input = st.text_input("Gemini API Key (Secrets 미설정 시 입력)", type="password")
    else:
        api_key_input = None
        st.success("✅ Secrets API 키 연결됨")

    n = st.slider("문제 수", 1, 30, 5)
    types_selected = st.multiselect("문제 유형", ["객관식", "OX", "단답형"], ["객관식"])
    difficulty = st.select_slider("난이도", ["기초","보통","어려움","시험 수준"], value="보통")
    styles = st.multiselect(
        "출제 스타일",
        ["핵심 개념", "개념 비교", "그림/도표 해석", "숫자/수치", "헷갈리는 선지", "응용"],
        ["핵심 개념", "그림/도표 해석"]
    )
    model = st.selectbox(
        "Gemini 모델",
        ["gemini-3.6-flash", "gemini-3.8-flash"],
        index=0
    )

uploaded = st.file_uploader("📎 PDF 또는 PPTX 업로드", type=["pdf","pptx"])

if uploaded:
    try:
        slides = extract(uploaded)
        st.success(f"{uploaded.name} 분석 준비 완료 — {len(slides)}개 페이지/슬라이드")
    except Exception as e:
        st.error(str(e))
        st.stop()

    tab1, tab2, tab3 = st.tabs(["📖 슬라이드 확인", "🧠 AI 요약/분석", "📝 생성된 문제"])

    with tab1:
        for s in slides:
            with st.expander(f"Slide/Page {s['page']}"):
                st.write(s.get("text","") or "(텍스트 없음)")
                if s.get("images"):
                    st.caption(f"추출된 시각자료 {len(s['images'])}개")

    with tab2:
        if st.button("🔎 분석 및 문제 생성 시작", type="primary"):
            settings = {
                "n": n, 
                "types": ", ".join(types_selected),
                "difficulty": difficulty,
                "styles": ", ".join(styles),
                "model": model,
                "api_key_input": api_key_input
            }
            with st.spinner("AI가 분석 중입니다... (서버 상태에 따라 15~30초 소요될 수 있습니다)"):
                result, err = ai_analyze(slides, settings)
            if err:
                st.error(err)
            else:
                st.session_state["result"] = result

        if "result" in st.session_state:
            r = st.session_state["result"]
            st.subheader("📌 핵심 요약")
            st.write(r.get("study_summary",""))
            
            st.subheader("💡 핵심 개념")
            for x in r.get("key_concepts", []):
                st.markdown(f"**{x.get('concept','')}** (중요도: {x.get('importance','')})")
                st.write(x.get("explanation",""))
                
            st.subheader("📊 시각자료 시험 포인트")
            for x in r.get("visual_points", []):
                st.markdown(f"**Slide {x.get('slide','?')} — {x.get('type','')}**")
                st.write(x.get("what_it_shows",""))
                st.caption(f"시험 포인트: {x.get('exam_point','')}")

    with tab3:
        if "result" not in st.session_state:
            st.info("먼저 'AI 요약/분석' 탭에서 분석을 진행해주세요.")
        else:
            qs = st.session_state["result"].get("questions", [])
            if not qs:
                st.warning("생성된 문제가 없습니다.")
            for i, q in enumerate(qs, 1):
                st.markdown(f"### Q{i}. {q.get('question','')}")
                st.caption(f"출처 Slide {q.get('source_slide','?')} | 유형: {q.get('type','')}")
                for c in q.get("choices", []):
                    st.write(c)
                with st.expander("정답 및 해설 보기"):
                    st.write("**정답:**", q.get("answer",""))
                    st.write("**해설:**", q.get("explanation",""))

else:
    st.info("PDF 또는 PPTX 파일을 업로드해 주세요.")