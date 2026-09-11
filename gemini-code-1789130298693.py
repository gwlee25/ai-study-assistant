import io, json, os, base64
from pathlib import Path

import streamlit as st
from pypdf import PdfReader
from pptx import Presentation
from google import genai
from google.genai import types

st.set_page_config(page_title="AI 학습 도우미 v2", page_icon="📚", layout="wide")
st.title("📚 AI 학습 도우미 v2")
st.caption("텍스트 + 그림/도표를 함께 분석하여 시험 대비 자료와 문제를 만드는 프로토타입")

# ---------- extraction ----------
def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append({"page": i, "text": page.extract_text() or ""})
    return pages

def pptx_to_images(data):
    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        images = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.shape_type == 13:  # picture
                try:
                    blob = shape.image.blob
                    mime = "image/png" if shape.image.ext in ("png",) else "image/jpeg"
                    images.append({
                        "mime": mime,
                        "raw_bytes": blob
                    })
                except Exception:
                    pass
        slides.append({"page": i, "text": "\n".join(texts), "images": images})
    return slides

def extract(uploaded):
    data = uploaded.getvalue()
    ext = Path(uploaded.name).suffix.lower()
    if ext == ".pdf":
        pages = extract_pdf(data)
        for p in pages:
            p["images"] = []
        return pages
    if ext == ".pptx":
        return pptx_to_images(data)
    raise ValueError("PDF 또는 PPTX만 지원합니다.")

# ---------- Gemini ----------
def ai_analyze(slides, settings):
    # 환경변수 또는 streamlit secrets 우선 확인
    key = os.getenv("GEMINI_API_KEY") or st.secrets.get("GEMINI_API_KEY", None)
    if not key:
        return None, "GEMINI_API_KEY가 설정되지 않았습니다. 사이드바에 입력하거나 환경변수를 설정해주세요."

    client = genai.Client(api_key=key)

    system_instruction = f"""
너는 대학 강의자료 기반 시험 대비 학습 도우미다.

반드시 제공된 강의자료를 근거로 분석한다.
자료에 없는 사실을 임의로 추가하지 않는다.
슬라이드의 그림, 현미경 이미지, 그래프, 표, 도식, 화살표, 라벨, 축, 범례 등도
가능한 경우 적극적으로 읽고 시험 포인트로 추출한다.

[사용자 설정]
문제 수: {settings['n']}
문제 유형: {settings['types']}
난이도: {settings['difficulty']}
출제 스타일: {settings['styles']}

출력은 반드시 순수한 JSON 형식으로만 응답해야 한다.
JSON 스키마:
{{
  "study_summary": "전체 요약",
  "key_concepts": [
    {{"concept":"개념명", "explanation":"설명", "importance":"high/medium/low"}}
  ],
  "visual_points": [
    {{"slide":1, "type":"diagram/image/graph/table", "what_it_shows":"내용", "exam_point":"시험 포인트"}}
  ],
  "questions": [
    {{
      "type":"객관식/OX/단답형",
      "source_slide": 1,
      "uses_visual": true,
      "question":"문제 내용",
      "choices":["A. ...","B. ...","C. ...","D. ..."],
      "answer":"정답",
      "explanation":"해설"
    }}
  ]
}}
"""

    contents = []
    # 슬라이드별 텍스트와 추출된 이미지를 순서대로 content 목록에 추가
    for s in slides:
        txt = s.get("text", "")
        slide_text_block = f"\n\n--- SLIDE/PAGE {s['page']} ---\n{txt[:10000]}"
        contents.append(slide_text_block)

        for img in s.get("images", []):
            contents.append(
                types.Part.from_bytes(
                    data=img["raw_bytes"],
                    mime_type=img["mime"]
                )
            )

    try:
        response = client.models.generate_content(
            model=settings["model"],
            contents=contents,
            config=types.GenerateContentConfig(
                system_instruction=system_instruction,
                response_mime_type="application/json",
                temperature=0.2
            )
        )
        return json.loads(response.text), None
    except Exception as e:
        return None, str(e)

# ---------- UI ----------
with st.sidebar:
    st.header("⚙️ API 및 문제 설정")
    api_key_input = st.text_input("Gemini API Key", type="password", placeholder="AIzaSy...")
    if api_key_input:
        os.environ["GEMINI_API_KEY"] = api_key_input

    n = st.slider("문제 수", 1, 30, 10)
    types_selected = st.multiselect("문제 유형", ["객관식", "OX", "단답형"], ["객관식"])
    difficulty = st.select_slider("난이도", ["기초","보통","어려움","시험 수준"], value="보통")
    styles = st.multiselect(
        "출제 스타일",
        ["핵심 개념", "개념 비교", "그림/도표 해석", "숫자/수치", "헷갈리는 선지", "응용"],
        ["핵심 개념", "개념 비교", "그림/도표 해석"]
    )
    model = st.selectbox(
        "Gemini 모델",
        ["gemini-2.5-flash", "gemini-2.5-pro"],
        index=0,
        help="flash 모델이 속도가 빠르고 무료 사용량 한도가 넉넉합니다."
    )

uploaded = st.file_uploader("📎 PDF 또는 PPTX 업로드", type=["pdf","pptx"])

if uploaded:
    try:
        slides = extract(uploaded)
        st.success(f"{uploaded.name} 분석 준비 완료 — {len(slides)}개 페이지/슬라이드")
    except Exception as e:
        st.error(str(e))
        st.stop()

    tab1, tab2, tab3 = st.tabs(["📖 자료 구조", "🧠 AI 분석", "📝 퀴즈 확인"])

    with tab1:
        for s in slides:
            with st.expander(f"Slide/Page {s['page']}"):
                st.write(s.get("text","") or "(텍스트 없음)")
                if s.get("images"):
                    st.caption(f"추출된 이미지 {len(s['images'])}개")

    with tab2:
        if st.button("🔎 텍스트 + 시각자료 분석", type="primary"):
            settings = {
                "n": n, 
                "types": ", ".join(types_selected),
                "difficulty": difficulty,
                "styles": ", ".join(styles),
                "model": model
            }
            with st.spinner("Gemini가 텍스트와 그림/도표를 분석하고 문제를 생성하는 중..."):
                result, err = ai_analyze(slides, settings)
            if err:
                st.error(err)
            else:
                st.session_state["result"] = result

        if "result" in st.session_state:
            r = st.session_state["result"]
            st.subheader("핵심 요약")
            st.write(r.get("study_summary",""))
            st.subheader("핵심 개념")
            for x in r.get("key_concepts", []):
                st.markdown(f"**{x['concept']}** · 중요도: {x['importance']}")
                st.write(x["explanation"])
            st.subheader("시각자료 시험 포인트")
            for x in r.get("visual_points", []):
                st.markdown(f"**Slide {x['slide']} — {x['type']}**")
                st.write(x["what_it_shows"])
                st.write("시험 포인트:", x["exam_point"])

    with tab3:
        if "result" not in st.session_state:
            st.info("먼저 'AI 분석' 탭에서 분석을 실행하세요.")
        else:
            qs = st.session_state["result"].get("questions", [])
            if not qs:
                st.warning("생성된 문제가 없습니다.")
            for i, q in enumerate(qs, 1):
                st.markdown(f"### Q{i}. {q.get('question','')}")
                st.caption(f"출처 Slide {q.get('source_slide','?')} | 유형: {q.get('type','')}")
                for c in q.get("choices", []):
                    st.write(c)
                with st.expander("정답 및 해설"):
                    st.write("정답:", q.get("answer",""))
                    st.write("해설:", q.get("explanation",""))

else:
    st.info("자료를 업로드하면 텍스트뿐 아니라 PPT에 포함된 그림/도표도 문제 출제 대상으로 사용할 수 있습니다.")