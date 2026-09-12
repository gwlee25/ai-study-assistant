import io, json, os, time, base64
from datetime import datetime
from pathlib import Path

import streamlit as st
from pypdf import PdfReader
from pptx import Presentation
from google import genai
from google.genai import types
from supabase import create_client, Client

st.set_page_config(page_title="AI 강의 노트 & 퀴즈 갤러리", page_icon="📚", layout="wide")

# ---------- 1. Supabase 초기화 ----------
raw_url = st.secrets.get("SUPABASE_URL") or os.getenv("SUPABASE_URL") or ""
supabase_key = st.secrets.get("SUPABASE_KEY") or os.getenv("SUPABASE_KEY") or ""
supabase_url = raw_url.strip().rstrip("/")

if not supabase_url or not supabase_key:
    st.error("⚠️ Supabase 연동 정보(SUPABASE_URL, SUPABASE_KEY)가 Secrets에 설정되지 않았습니다.")
    st.stop()

@st.cache_resource
def init_supabase(url: str, key: str) -> Client:
    return create_client(url, key)

supabase = init_supabase(supabase_url, supabase_key)

# ---------- 2. 파일 파싱 및 시각자료 추출 ----------
def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append({"page": i, "text": page.extract_text() or "", "images": []})
    return pages

def pptx_to_images(data):
    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        images = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.shape_type == 13:  # picture
                try:
                    blob = shape.image.blob
                    mime = "image/png" if shape.image.ext in ("png",) else "image/jpeg"
                    images.append({
                        "mime": mime,
                        "b64": base64.b64encode(blob).decode("utf-8")
                    })
                except Exception:
                    pass
        slides.append({"page": i, "text": "\n".join(texts), "images": images})
    return slides

def extract(data, filename):
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".pptx":
        return pptx_to_images(data)
    raise ValueError("PDF 또는 PPTX 파일만 지원합니다.")

# ---------- 3. DB 함수 ----------
def fetch_user_folders(username):
    try:
        res = supabase.table("user_folders").select("folder_name").eq("username", username).order("created_at").execute()
        return [r["folder_name"] for r in (res.data or [])]
    except Exception as e:
        st.error(f"폴더 목록 조회 오류: {e}")
        return []

def create_folder_in_db(username, folder_name):
    try:
        supabase.table("user_folders").upsert({
            "username": username,
            "folder_name": folder_name
        }, on_conflict="username,folder_name").execute()
        return True
    except Exception as e:
        st.error(f"폴더 생성 오류: {e}")
        return False

def fetch_user_data(username):
    try:
        response = supabase.table("user_documents").select("*").eq("username", username).execute()
        return response.data or []
    except Exception as e:
        st.error(f"데이터베이스 조회 오류: {e}")
        return []

def save_document(username, folder_name, filename, slides):
    existing = supabase.table("user_documents").select("id")\
        .eq("username", username)\
        .eq("folder_name", folder_name)\
        .eq("filename", filename).execute()
    
    if existing.data:
        supabase.table("user_documents").update({"slides_data": slides, "analysis_result": None})\
            .eq("id", existing.data[0]["id"]).execute()
    else:
        supabase.table("user_documents").insert({
            "username": username,
            "folder_name": folder_name,
            "filename": filename,
            "slides_data": slides,
            "analysis_result": None
        }).execute()

def update_analysis_result(doc_id, result):
    supabase.table("user_documents").update({"analysis_result": result}).eq("id", doc_id).execute()

def delete_document(doc_id):
    supabase.table("user_documents").delete().eq("id", doc_id).execute()

def delete_folder(username, folder_name):
    supabase.table("user_documents").delete().eq("username", username).eq("folder_name", folder_name).execute()
    supabase.table("user_folders").delete().eq("username", username).eq("folder_name", folder_name).execute()

# ---------- 4. 삭제 확인 모달 팝업 ----------
@st.dialog("⚠️ 폴더 삭제 확인")
def confirm_delete_folder_dialog(username, folder_name):
    st.write(f"정말로 **'{folder_name}'** 폴더를 삭제하시겠습니까?")
    st.warning("폴더와 그 안에 보관된 모든 문서 및 데이터가 함께 영구 삭제됩니다.")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("예, 삭제합니다", type="primary", use_container_width=True):
            delete_folder(username, folder_name)
            st.rerun()
    with c2:
        if st.button("취소", use_container_width=True):
            st.rerun()

@st.dialog("⚠️ 파일 삭제 확인")
def confirm_delete_file_dialog(doc_id, filename):
    st.write(f"정말로 **'{filename}'** 파일을 삭제하시겠습니까?")
    st.warning("클라우드에 저장된 슬라이드 데이터 및 분석 결과가 영구 삭제됩니다.")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("예, 삭제합니다", type="primary", use_container_width=True):
            delete_document(doc_id)
            st.rerun()
    with c2:
        if st.button("취소", use_container_width=True):
            st.rerun()

# ---------- 5. Gemini AI 분리 분석 로직 ----------
def call_gemini_api(contents, system_instruction, model_name):
    api_key = st.secrets.get("GEMINI_API_KEY") or os.getenv("GEMINI_API_KEY")
    if not api_key:
        return None, "Gemini API 키가 설정되지 않았습니다."

    client = genai.Client(api_key=api_key)
    models_to_try = [model_name, "gemini-3.6-flash", "gemini-3.8-flash"]
    models_to_try = list(dict.fromkeys(models_to_try))

    last_error = ""
    for target_model in models_to_try:
        for _ in range(2):
            try:
                response = client.models.generate_content(
                    model=target_model,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        system_instruction=system_instruction,
                        response_mime_type="application/json",
                        temperature=0.2
                    )
                )
                return json.loads(response.text), None
            except Exception as e:
                last_error = str(e)
                if "503" in last_error:
                    time.sleep(3)
                    continue
                elif "404" in last_error:
                    break
                else:
                    return None, f"오류 발생: {last_error}"

    return None, f"서버 과부하로 실패했습니다. 잠시 후 다시 시도해주세요. ({last_error})"

def prepare_gemini_contents(slides):
    contents = []
    for s in slides:
        txt = s.get("text", "")
        contents.append(f"\n\n--- SLIDE/PAGE {s['page']} ---\n{txt[:10000]}")
        for img in s.get("images", []):
            try:
                raw_bytes = base64.b64decode(img["b64"])
                contents.append(types.Part.from_bytes(data=raw_bytes, mime_type=img["mime"]))
            except Exception:
                pass
    return contents

# A. 상세 한국어 요약 전용 호출
def ai_generate_summary(slides, model_name):
    instruction = """
You are an elite university lecture specialist.
Analyze the provided slides thoroughly and generate detailed Korean study notes.

[SUMMARY RULES - KOREAN]
1. 모든 슬라이드/페이지에 대해 내용을 생략하지 말고 누락 없이 매우 상세하게 분석하라.
2. 'page_summaries' 항목에 각 페이지 번호(page_num)별 핵심 주제(topic), 주요 개념 및 세부 내용(details)을 상세한 한국어로 작성하라.
3. 전공/학술 용어, 공식, 약어, 핵심 키워드는 원문 영어 용어를 반드시 병기하라 (예: 활동 전위(Action Potential)).
4. 슬라이드에 다이어그램, 도표, 이미지가 포함되어 있다면 시각자료 내용(visual_analysis)도 상세히 서술하라.

Output MUST strictly follow this JSON schema:
{
  "overall_summary": "전체 요약 (한국어)",
  "page_summaries": [
    {
      "page_num": 1,
      "topic": "슬라이드 주제",
      "details": ["세부 설명 1", "세부 설명 2"],
      "visual_analysis": "도표/그림 분석 (없으면 빈칸)"
    }
  ],
  "key_terminology": [
    {"term_en": "Term", "term_kr": "한국어 용어", "definition": "설명"}
  ]
}
"""
    contents = prepare_gemini_contents(slides)
    return call_gemini_api(contents, instruction, model_name)

# B. 영문 퀴즈(한국어 해설) 단독 생성 호출
def ai_generate_quiz(slides, settings):
    instruction = f"""
You are an elite university exam tutor.
Based on the provided slides, create rigorous exam questions.

[QUIZ RULES - QUESTIONS IN ENGLISH, EXPLANATION IN KOREAN]
1. Questions, choices, and answers MUST BE WRITTEN IN ENGLISH.
2. The 'explanation' field MUST BE WRITTEN IN DETAILED KOREAN (한국어 해설 필수).
3. Number of questions: {settings['n']}
4. Question types: {settings['types']}
5. Difficulty: {settings['difficulty']}

Output MUST strictly follow this JSON schema:
{{
  "questions": [
    {{
      "id": 1,
      "type": "Type",
      "source_page": 1,
      "question": "Question in English",
      "choices": ["A. ...", "B. ...", "C. ...", "D. ..."],
      "answer": "Answer in English",
      "explanation": "상세한 한국어 해설"
    }}
  ]
}}
"""
    contents = prepare_gemini_contents(slides)
    return call_gemini_api(contents, instruction, settings["model"])

# ---------- 6. 세션 네비게이션 상태 ----------
if "nav_view" not in st.session_state:
    st.session_state["nav_view"] = "folder"
if "active_folder" not in st.session_state:
    st.session_state["active_folder"] = None
if "active_doc_id" not in st.session_state:
    st.session_state["active_doc_id"] = None

# ---------- 7. 사이드바 UI ----------
with st.sidebar:
    st.header("👤 계정 접속")
    username_input = st.text_input("사용자 이름 / 핀번호", placeholder="예: user1234").strip()
    
    if not username_input:
        st.warning("계정 이름을 입력해야 보관함이 활성화됩니다.")
        st.stop()
        
    st.success(f"접속 중: **{username_input}**")

    st.markdown("---")
    st.header("⚙️ 퀴즈 생성 옵션")
    quiz_n = st.slider("문제 수", 1, 20, value=10)
    quiz_types = st.multiselect("문제 유형", ["Multiple Choice", "True/False", "Short Answer"], ["Multiple Choice"])
    quiz_diff = st.select_slider("난이도", ["Basic", "Intermediate", "Advanced", "Exam Level"], value="Intermediate")
    quiz_model = st.selectbox("Gemini 모델", ["gemini-3.6-flash", "gemini-3.8-flash"], index=0)

db_folders = fetch_user_folders(username_input)
user_records = fetch_user_data(username_input)

for r in user_records:
    fn = r.get("folder_name")
    if fn and fn not in db_folders:
        create_folder_in_db(username_input, fn)
        db_folders.append(fn)

if not db_folders:
    create_folder_in_db(username_input, "기본 강의자료")
    db_folders = ["기본 강의자료"]

# =========================================================
# 화면 1: 폴더 갤러리
# =========================================================
if st.session_state["nav_view"] == "folder":
    st.title("📂 내 강의자료 폴더")
    st.caption("폴더를 클릭하여 내부 PPT/PDF 목록을 확인하세요.")

    col_nf1, col_nf2 = st.columns([3, 1])
    with col_nf1:
        new_folder_val = st.text_input("새 폴더 생성", placeholder="새 폴더명 입력", label_visibility="collapsed")
    with col_nf2:
        if st.button("➕ 폴더 추가", use_container_width=True) and new_folder_val.strip():
            c_name = new_folder_val.strip()
            if c_name not in db_folders:
                create_folder_in_db(username_input, c_name)
                st.rerun()

    st.write("")

    cols_per_row = 3
    for i in range(0, len(db_folders), cols_per_row):
        cols = st.columns(cols_per_row)
        for j in range(cols_per_row):
            idx = i + j
            if idx < len(db_folders):
                f_name = db_folders[idx]
                f_docs = [r for r in user_records if r.get("folder_name") == f_name]
                file_count = len(f_docs)
                
                with cols[j]:
                    with st.container(border=True):
                        st.markdown(f"### 📁 {f_name}")
                        st.caption(f"보관된 자료: {file_count}개")
                        
                        b_col1, b_col2 = st.columns([2, 1])
                        with b_col1:
                            if st.button("📂 열기", key=f"open_f_{f_name}", use_container_width=True):
                                st.session_state["active_folder"] = f_name
                                st.session_state["nav_view"] = "file"
                                st.rerun()
                        with b_col2:
                            if st.button("🗑️", key=f"del_f_{f_name}", use_container_width=True, help="폴더 삭제"):
                                confirm_delete_folder_dialog(username_input, f_name)

# =========================================================
# 화면 2: 파일 갤러리
# =========================================================
elif st.session_state["nav_view"] == "file":
    curr_f = st.session_state["active_folder"]
    
    col_nav, _ = st.columns([1, 4])
    with col_nav:
        if st.button("⬅️ 전체 폴더 목록으로"):
            st.session_state["nav_view"] = "folder"
            st.session_state["active_folder"] = None
            st.rerun()

    st.title(f"📂 {curr_f}")
    st.caption("자료를 추가하거나, 카드를 클릭하여 AI 요약 및 퀴즈를 확인하세요.")

    with st.expander("📎 이 폴더에 새 강의자료 업로드 (PPTX, PDF)", expanded=False):
        uploaded_files = st.file_uploader("파일 선택 (복수 업로드 가능)", type=["pdf", "pptx"], accept_multiple_files=True)
        if uploaded_files:
            with st.spinner("파일 및 시각자료를 저장하는 중..."):
                for up in uploaded_files:
                    try:
                        extracted_slides = extract(up.getvalue(), up.name)
                        save_document(username_input, curr_f, up.name, extracted_slides)
                    except Exception as e:
                        st.error(f"{up.name} 처리 오류: {e}")
                st.success("클라우드 저장소에 보관되었습니다!")
                time.sleep(1)
                st.rerun()

    st.markdown("---")
    
    f_docs = [r for r in user_records if r.get("folder_name") == curr_f]
    
    if not f_docs:
        st.info("이 폴더에 저장된 파일이 없습니다. 상단에서 자료를 업로드해 주세요.")
    else:
        cols_per_row = 3
        for i in range(0, len(f_docs), cols_per_row):
            cols = st.columns(cols_per_row)
            for j in range(cols_per_row):
                idx = i + j
                if idx < len(f_docs):
                    doc = f_docs[idx]
                    fname = doc["filename"]
                    doc_id = doc["id"]
                    analysis_data = doc.get("analysis_result") or {}
                    has_summary = bool(analysis_data.get("overall_summary"))
                    quiz_count = len(analysis_data.get("quiz_history", []))
                    if not quiz_count and analysis_data.get("questions"):
                        quiz_count = 1
                    slide_count = len(doc.get("slides_data", []))
                    
                    with cols[j]:
                        with st.container(border=True):
                            icon = "📊" if fname.endswith(".pptx") else "📄"
                            st.markdown(f"### {icon} {fname}")
                            st.write(f"슬라이드/페이지: **{slide_count}장**")
                            
                            status_txt = []
                            if has_summary:
                                status_txt.append("요약 완료")
                            if quiz_count > 0:
                                status_txt.append(f"퀴즈 {quiz_count}세트")
                            
                            if status_txt:
                                st.success("✅ " + " | ".join(status_txt))
                            else:
                                st.info("⏳ 미분석 상태")
                                
                            c_open, c_del = st.columns([3, 1])
                            with c_open:
                                if st.button("📖 열기 및 분석", key=f"open_doc_{doc_id}", use_container_width=True):
                                    st.session_state["active_doc_id"] = doc_id
                                    st.session_state["nav_view"] = "detail"
                                    st.rerun()
                            with c_del:
                                if st.button("🗑️", key=f"del_doc_{doc_id}", use_container_width=True, help="파일 삭제"):
                                    confirm_delete_file_dialog(doc_id, fname)

# =========================================================
# 화면 3: 문서 상세 뷰
# =========================================================
elif st.session_state["nav_view"] == "detail":
    doc_id = st.session_state["active_doc_id"]
    current_doc = next((r for r in user_records if r["id"] == doc_id), None)
    
    if not current_doc:
        st.warning("문서를 찾을 수 없습니다.")
        if st.button("⬅️ 목록으로 돌아가기"):
            st.session_state["nav_view"] = "file"
            st.rerun()
        st.stop()

    curr_f = current_doc["folder_name"]
    fname = current_doc["filename"]
    slides = current_doc["slides_data"]
    
    # 분석 데이터 구조 정규화 (하위 호환)
    analysis_result = current_doc.get("analysis_result") or {}
    if "quiz_history" not in analysis_result:
        analysis_result["quiz_history"] = []
        # 과거 단일 questions 필드가 있었다면 1회차로 자동 변환
        if "questions" in analysis_result and analysis_result["questions"]:
            analysis_result["quiz_history"].append({
                "set_name": "Quiz Set #1 (기존)",
                "created_at": "이전 생성분",
                "questions": analysis_result["questions"]
            })

    col_back, _ = st.columns([1, 4])
    with col_back:
        if st.button(f"⬅️ '{curr_f}' 파일 목록으로"):
            st.session_state["nav_view"] = "file"
            st.session_state["active_doc_id"] = None
            st.rerun()

    st.title(f"📖 {fname}")
    st.caption(f"폴더: {curr_f} | 슬라이드 수: {len(slides)}장")

    tab1, tab2, tab3 = st.tabs(["📑 슬라이드 원본", "📝 상세 한국어 요약노트", "🎯 English Quiz (히스토리 보관)"])

    # --- 탭 1: 슬라이드 원본 ---
    with tab1:
        st.info(f"총 {len(slides)}개의 슬라이드/페이지가 보관되어 있습니다.")
        for s in slides:
            with st.expander(f"Slide/Page {s['page']}"):
                st.write(s.get("text") or "(텍스트 없음)")
                if s.get("images"):
                    st.caption(f"📸 포함된 원본 시각자료: {len(s['images'])}개")

    # --- 탭 2: 요약노트 단독 생성 및 열람 ---
    with tab2:
        col_btn2, _ = st.columns([1, 3])
        with col_btn2:
            run_summary_btn = st.button("🚀 상세 요약노트 생성/갱신", type="primary", key="btn_run_summary")

        if run_summary_btn:
            with st.spinner("AI가 페이지별 세부 내용과 시각자료를 빠짐없이 요약 중입니다..."):
                summary_data, err = ai_generate_summary(slides, quiz_model)
            if err:
                st.error(err)
            else:
                # 기존 퀴즈 히스토리는 유지하고 요약 부분만 갱신
                analysis_result["overall_summary"] = summary_data.get("overall_summary")
                analysis_result["page_summaries"] = summary_data.get("page_summaries")
                analysis_result["key_terminology"] = summary_data.get("key_terminology")
                update_analysis_result(doc_id, analysis_result)
                st.success("상세 요약노트가 생성 및 저장되었습니다!")
                st.rerun()

        if analysis_result.get("overall_summary"):
            st.markdown("### 📌 전반적인 강의 개요")
            st.info(analysis_result.get("overall_summary", "요약 내용이 없습니다."))

            if analysis_result.get("key_terminology"):
                st.markdown("### 📖 핵심 학술/전문 용어 (Terminology)")
                cols = st.columns(2)
                for idx, term in enumerate(analysis_result.get("key_terminology", [])):
                    with cols[idx % 2]:
                        st.markdown(f"**🔹 {term.get('term_en','')}** ({term.get('term_kr','')})")
                        st.write(term.get("definition",""))

            st.markdown("---")
            st.markdown("### 📑 슬라이드/페이지별 상세 정리")
            for page_data in analysis_result.get("page_summaries", []):
                with st.expander(f"📍 [Page {page_data.get('page_num')}] {page_data.get('topic')}", expanded=True):
                    st.markdown("**세부 내용:**")
                    for d in page_data.get("details", []):
                        st.markdown(f"- {d}")
                    if page_data.get("visual_analysis"):
                        st.markdown(f"🖼️ **시각자료 분석:** {page_data.get('visual_analysis')}")
        else:
            st.info("상단의 **'🚀 상세 요약노트 생성/갱신'** 버튼을 누르면 AI 요약이 실행됩니다.")

    # --- 탭 3: 퀴즈 단독 생성 및 회차별 히스토리 열람 ---
    with tab3:
        col_btn3, _ = st.columns([1, 3])
        with col_btn3:
            run_quiz_btn = st.button("🎲 새로운 퀴즈 세트 생성", type="primary", key="btn_run_quiz")

        if run_quiz_btn:
            settings = {
                "n": quiz_n,
                "types": ", ".join(quiz_types),
                "difficulty": quiz_diff,
                "model": quiz_model
            }
            with st.spinner("AI가 슬라이드 기반으로 새로운 영문 퀴즈를 출제 중입니다..."):
                quiz_data, err = ai_generate_quiz(slides, settings)
            if err:
                st.error(err)
            else:
                new_q_list = quiz_data.get("questions", [])
                now_str = datetime.now().strftime("%m/%d %H:%M")
                set_num = len(analysis_result.get("quiz_history", [])) + 1
                new_set = {
                    "set_name": f"Quiz Set #{set_num} ({now_str} 생성 - {len(new_q_list)}문항)",
                    "created_at": now_str,
                    "questions": new_q_list
                }
                # 맨 앞에 추가하여 최신 퀴즈가 기본으로 뜨도록 처리
                analysis_result["quiz_history"].insert(0, new_set)
                update_analysis_result(doc_id, analysis_result)
                st.success(f"새로운 {new_set['set_name']}이 추가되었습니다!")
                st.rerun()

        quiz_history = analysis_result.get("quiz_history", [])
        if not quiz_history:
            st.info("아직 생성된 퀴즈가 없습니다. 상단의 **'🎲 새로운 퀴즈 세트 생성'** 버튼을 눌러 문제를 만들어보세요.")
        else:
            st.markdown("---")
            # 저장된 회차 선택 드롭다운
            set_names = [q_set["set_name"] for q_set in quiz_history]
            selected_set_name = st.selectbox("📚 다시 볼 퀴즈 세트 선택", set_names, index=0)
            
            selected_set = next((q_set for q_set in quiz_history if q_set["set_name"] == selected_set_name), None)
            
            if selected_set:
                questions = selected_set.get("questions", [])
                st.markdown(f"### 📝 {selected_set['set_name']}")
                st.caption("Questions & Choices: English | Explanations: Korean")
                
                for idx, q in enumerate(questions, 1):
                    st.markdown(f"#### Q{idx}. {q.get('question')}")
                    st.caption(f"Source: Page {q.get('source_page')} | Type: {q.get('type')}")
                    for ch in q.get("choices", []):
                        st.write(ch)
                    with st.expander(f"정답 및 한국어 해설 확인 (Q{idx})"):
                        st.markdown(f"**Answer:** `{q.get('answer')}`")
                        st.markdown(f"**해설:** {q.get('explanation')}")
                    st.write("")
