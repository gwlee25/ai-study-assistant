
import io, json, os, base64
from pathlib import Path

import streamlit as st
from pypdf import PdfReader
from pptx import Presentation

st.set_page_config(page_title="AI 학습 도우미 v2", page_icon="📚", layout="wide")
st.title("📚 AI 학습 도우미 v2")
st.caption("텍스트 + 그림/도표를 함께 분석하여 시험 대비 자료와 문제를 만드는 프로토타입")

# ---------- extraction ----------
def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append({"page": i, "text": page.extract_text() or ""})
    return pages

def pptx_to_images(data):
    # PowerPoint itself does not render reliably in a lightweight Streamlit app.
    # We still extract slide text and embedded images. The API can receive those images.
    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        images = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.shape_type == 13:  # picture
                try:
                    blob = shape.image.blob
                    mime = "image/png" if shape.image.ext in ("png",) else "image/jpeg"
                    images.append({
                        "mime": mime,
                        "b64": base64.b64encode(blob).decode("utf-8")
                    })
                except Exception:
                    pass
        slides.append({"page": i, "text": "\n".join(texts), "images": images})
    return slides

def pdf_pages_as_images(data):
    # Optional high-quality visual analysis can be done by converting PDF pages
    # to images before sending to a vision-capable model.
    # This function intentionally avoids external binaries in the prototype.
    return []

def extract(uploaded):
    data = uploaded.getvalue()
    ext = Path(uploaded.name).suffix.lower()
    if ext == ".pdf":
        pages = extract_pdf(data)
        for p in pages:
            p["images"] = []
        return pages
    if ext == ".pptx":
        return pptx_to_images(data)
    raise ValueError("PDF 또는 PPTX만 지원합니다.")

# ---------- OpenAI ----------
def ai_analyze(slides, settings):
    key = os.getenv("OPENAI_API_KEY")
    if not key:
        return None, "OPENAI_API_KEY 환경변수가 없습니다."

    from openai import OpenAI
    client = OpenAI(api_key=key)

    content = [{
        "type": "input_text",
        "text": f"""
너는 대학 강의자료 기반 시험 대비 학습 도우미다.

반드시 제공된 강의자료를 근거로 분석한다.
자료에 없는 사실을 임의로 추가하지 않는다.
슬라이드의 그림, 현미경 이미지, 그래프, 표, 도식, 화살표, 라벨, 축, 범례 등도
가능한 경우 적극적으로 읽고 시험 포인트로 추출한다.

[사용자 설정]
문제 수: {settings['n']}
문제 유형: {settings['types']}
난이도: {settings['difficulty']}
출제 스타일: {settings['styles']}

각 시각자료에 대해 다음을 판단한다.
1. 무엇을 보여주는지
2. 그림의 핵심 관계/과정/비교
3. 라벨·축·범례·수치에서 중요한 정보
4. 텍스트 설명과 그림이 어떻게 연결되는지
5. 그림을 보고 물을 수 있는 시험 문제

출력은 반드시 JSON으로 한다.
{{
  "study_summary": "...",
  "key_concepts": [
    {{"concept":"...", "explanation":"...", "importance":"high/medium/low"}}
  ],
  "visual_points": [
    {{"slide":1, "type":"diagram/image/graph/table", "what_it_shows":"...", "exam_point":"..."}}
  ],
  "questions": [
    {{
      "type":"객관식/OX/단답형",
      "source_slide": 1,
      "uses_visual": true,
      "question":"...",
      "choices":["A. ...","B. ...","C. ...","D. ..."],
      "answer":"...",
      "explanation":"..."
    }}
  ]
}}

[강의자료]
"""
    }]

    for s in slides:
        txt = s.get("text", "")
        content[0]["text"] += f"\n\n--- SLIDE/PAGE {s['page']} ---\n{txt[:10000]}"
        for img in s.get("images", []):
            content.append({
                "type": "input_image",
                "image_url": f"data:{img['mime']};base64,{img['b64']}"
            })

    try:
        r = client.responses.create(
            model=settings["model"],
            input=[{"role": "user", "content": content}],
            instructions="정확하고 보수적으로 분석하고, JSON 이외의 텍스트를 출력하지 마라."
        )
        return json.loads(r.output_text), None
    except Exception as e:
        return None, str(e)

# ---------- UI ----------
with st.sidebar:
    st.header("⚙️ 문제 설정")
    n = st.slider("문제 수", 1, 30, 10)
    types = st.multiselect("문제 유형", ["객관식", "OX", "단답형"], ["객관식"])
    difficulty = st.select_slider("난이도", ["기초","보통","어려움","시험 수준"], value="보통")
    styles = st.multiselect(
        "출제 스타일",
        ["핵심 개념", "개념 비교", "그림/도표 해석", "숫자/수치", "헷갈리는 선지", "응용"],
        ["핵심 개념", "개념 비교", "그림/도표 해석"]
    )
    model = st.text_input("Vision 지원 OpenAI 모델", "gpt-5.6-luna")

uploaded = st.file_uploader("📎 PDF 또는 PPTX 업로드", type=["pdf","pptx"])

if uploaded:
    try:
        slides = extract(uploaded)
        st.success(f"{uploaded.name} 분석 준비 완료 — {len(slides)}개 페이지/슬라이드")
    except Exception as e:
        st.error(str(e))
        st.stop()

    tab1, tab2, tab3 = st.tabs(["📖 자료 구조", "🧠 AI 분석", "📝 그림/도표 문제"])

    with tab1:
        for s in slides:
            with st.expander(f"Slide/Page {s['page']}"):
                st.write(s.get("text","") or "(텍스트 없음)")
                if s.get("images"):
                    st.caption(f"추출된 이미지 {len(s['images'])}개")

    with tab2:
        if st.button("🔎 텍스트 + 시각자료 분석", type="primary"):
            settings = {
                "n": n, "types": ", ".join(types),
                "difficulty": difficulty,
                "styles": ", ".join(styles),
                "model": model
            }
            with st.spinner("텍스트와 그림/도표를 함께 분석하는 중..."):
                result, err = ai_analyze(slides, settings)
            if err:
                st.error(err)
                st.info("OPENAI_API_KEY를 설정한 뒤 실행하면 AI 분석이 가능합니다.")
            else:
                st.session_state["result"] = result

        if "result" in st.session_state:
            r = st.session_state["result"]
            st.subheader("핵심 요약")
            st.write(r.get("study_summary",""))
            st.subheader("핵심 개념")
            for x in r.get("key_concepts", []):
                st.markdown(f"**{x['concept']}** · 중요도: {x['importance']}")
                st.write(x["explanation"])
            st.subheader("시각자료 시험 포인트")
            for x in r.get("visual_points", []):
                st.markdown(f"**Slide {x['slide']} — {x['type']}**")
                st.write(x["what_it_shows"])
                st.write("시험 포인트:", x["exam_point"])

    with tab3:
        if "result" not in st.session_state:
            st.info("먼저 '텍스트 + 시각자료 분석'을 실행하세요.")
        else:
            qs = [q for q in st.session_state["result"].get("questions", []) if q.get("uses_visual")]
            if not qs:
                st.warning("시각자료를 직접 사용하는 문제가 생성되지 않았습니다.")
            for i, q in enumerate(qs, 1):
                st.markdown(f"### Q{i}. {q['question']}")
                st.caption(f"출처 Slide {q.get('source_slide','?')}")
                for c in q.get("choices", []):
                    st.write(c)
                with st.expander("정답 및 해설"):
                    st.write("정답:", q.get("answer",""))
                    st.write(q.get("explanation",""))

else:
    st.info("자료를 업로드하면 텍스트뿐 아니라 PPT에 포함된 그림/도표도 문제 출제 대상으로 사용할 수 있습니다.")
